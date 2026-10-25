#!/usr/bin/env python3
"""Render a tailored resume JSON to an ATS-friendly .docx and .pdf.

Usage: python3 render_resume.py resume.json OUT_BASENAME
JSON keys: name, headline, contact (list of str), summary (str),
achievements (list), experience (list of {title, dates, company, location, bullets}),
skills (list of {label, items}), education (list of str), certifications (str).
Use **double asterisks** inside any text for bold.
"""
import json, re, subprocess, sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1F, 0x5F, 0xBF)
GREY = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"
BODY = 10.5


def runs(p, text, size=BODY, color=None, italic=False, bold_all=False):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.font.name = FONT
        r.bold = bold_all or (i % 2 == 1)
        r.italic = italic
        if color:
            r.font.color.rgb = color


def spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


def rule(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F5FBF")):
        bot.set(qn("w:" + k), v)
    b.append(bot)
    pPr.append(b)


def heading(doc, text):
    p = doc.add_paragraph()
    spacing(p, 8, 3)
    runs(p, text.upper(), size=11.5, color=BLUE, bold_all=True)
    rule(p)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    spacing(p, 0, 1.5, 1.0)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    runs(p, text)


def build(data, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(BODY)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 0, 0)
    runs(p, data["name"].upper(), size=20, bold_all=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 1, 0)
    runs(p, data["headline"], size=10.5, color=BLUE, bold_all=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 1, 2)
    runs(p, "  |  ".join(data["contact"]), size=9.5, color=GREY)

    heading(doc, "Summary")
    p = doc.add_paragraph(); spacing(p, 0, 0, 1.05); runs(p, data["summary"])

    if data.get("achievements"):
        heading(doc, "Key Achievements")
        for a in data["achievements"]:
            bullet(doc, a)

    heading(doc, "Experience")
    for job in data["experience"]:
        p = doc.add_paragraph(); spacing(p, 5, 0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
        runs(p, job["title"], size=11, bold_all=True)
        runs(p, "\t" + job["dates"], size=10, bold_all=True)
        p = doc.add_paragraph(); spacing(p, 0, 2)
        runs(p, f'{job["company"]}  |  {job["location"]}', size=10, color=BLUE, italic=True)
        for b in job["bullets"]:
            bullet(doc, b)

    if data.get("skills"):
        heading(doc, "Tools & Platforms")
        for s in data["skills"]:
            p = doc.add_paragraph(); spacing(p, 0, 1.5)
            runs(p, f'**{s["label"]}:** ' + " • ".join(s["items"]))

    heading(doc, "Education")
    for e in data["education"]:
        p = doc.add_paragraph(); spacing(p, 0, 1); runs(p, e)
    if data.get("certifications"):
        heading(doc, "Certifications")
        p = doc.add_paragraph(); spacing(p, 0, 0); runs(p, data["certifications"])

    doc.save(out + ".docx")
    outdir = os.path.dirname(os.path.abspath(out))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, out + ".docx"],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pages = subprocess.run(["pdfinfo", out + ".pdf"], capture_output=True, text=True).stdout
    n = int(re.search(r"Pages:\s+(\d+)", pages).group(1))
    print(f"{out}.docx / .pdf -> {n} page(s)")
    if n > 2:
        print("WARNING: over 2 pages - trim bullets and re-render")


if __name__ == "__main__":
    build(json.load(open(sys.argv[1])), sys.argv[2])
