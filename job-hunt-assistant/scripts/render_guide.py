#!/usr/bin/env python3
"""Render an interview guide JSON to .docx and .pdf (same look as the resumes).

Usage: python3 render_guide.py guide.json OUT_BASENAME
JSON keys:
  company, title, candidate, stage (e.g. "Recruiter pre-screen"), date (YYYY-MM-DD or ""), interviewer ("" if unknown),
  snapshot: [str]                      company facts, each ending with its source in brackets
  pitch: str                           60-second "tell me about yourself"
  why_them: str                        why this company and this role
  story_map: [{story, use_for}]        which story-bank stories to rehearse, and for what
  behavioural: [{q, story, s, t, a, r, tip}]
  role_questions: [{q, answer}]
  gaps: [{q, answer}]
  plan: {d30: [str], d60: [str], d90: [str]}
  ask_them: [str]
  logistics: [str]                     salary line, format, what to have open, follow-up timing
Use **double asterisks** for bold.
"""
import json, re, subprocess, sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1F, 0x5F, 0xBF)
GREY = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"
BODY = 10.5


def runs(p, text, size=BODY, color=None, italic=False, bold_all=False):
    for i, part in enumerate(re.split(r"\*\*", str(text))):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.font.name = FONT
        r.bold = bold_all or (i % 2 == 1)
        r.italic = italic
        if color:
            r.font.color.rgb = color


def spacing(p, before=0, after=0, line=1.05):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


def rule(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F5FBF")):
        bot.set(qn("w:" + k), v)
    b.append(bot)
    pPr.append(b)


def shade(p, fill="EAF0FA"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def heading(doc, text, keep=True):
    p = doc.add_paragraph()
    spacing(p, 10, 4)
    p.paragraph_format.keep_with_next = keep
    runs(p, text.upper(), size=11.5, color=BLUE, bold_all=True)
    rule(p)


def para(doc, text, after=4, **kw):
    p = doc.add_paragraph(); spacing(p, 0, after); runs(p, text, **kw); return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    spacing(p, 0, 2)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    runs(p, text)


def qhead(doc, n, q):
    p = doc.add_paragraph(); spacing(p, 6, 2)
    p.paragraph_format.keep_with_next = True
    runs(p, f"{n}. {q}", size=11, bold_all=True)


def build(g, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.7))
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(BODY)

    p = doc.add_paragraph(); spacing(p, 0, 0)
    runs(p, "INTERVIEW GUIDE", size=10, color=GREY, bold_all=True)
    p = doc.add_paragraph(); spacing(p, 2, 0)
    runs(p, f'{g["title"]} · {g["company"]}', size=18, bold_all=True)
    meta = [x for x in (g.get("stage"), g.get("date"), (f'with {g["interviewer"]}' if g.get("interviewer") else None)) if x]
    p = doc.add_paragraph(); spacing(p, 2, 4)
    runs(p, "  |  ".join(meta + ([g["candidate"]] if g.get("candidate") else [])), size=10, color=GREY)

    if g.get("snapshot"):
        heading(doc, "Company snapshot")
        for s in g["snapshot"]:
            bullet(doc, s)

    heading(doc, "Tell me about yourself (60 seconds)")
    p = para(doc, g["pitch"], after=2); shade(p)
    if g.get("why_them"):
        heading(doc, "Why this company, why this role")
        para(doc, g["why_them"])

    if g.get("story_map"):
        heading(doc, "Stories to have ready")
        for s in g["story_map"]:
            bullet(doc, f'**{s["story"]}**: {s["use_for"]}')

    heading(doc, "Behavioural questions (STAR)")
    for i, b in enumerate(g["behavioural"], 1):
        qhead(doc, i, b["q"])
        if b.get("story"):
            para(doc, f'Story: {b["story"]}', after=1, size=9.5, color=GREY, italic=True)
        for k, lab in (("s", "Situation"), ("t", "Task"), ("a", "Action"), ("r", "Result")):
            if b.get(k):
                para(doc, f"**{lab}:** {b[k]}", after=1)
        if b.get("tip"):
            para(doc, f'Tip: {b["tip"]}', after=3, size=9.5, color=GREY, italic=True)

    if g.get("role_questions"):
        heading(doc, "Role and skills questions")
        for i, b in enumerate(g["role_questions"], 1):
            qhead(doc, i, b["q"]); para(doc, b["answer"])

    if g.get("gaps"):
        heading(doc, "Gap questions: honest answers")
        for i, b in enumerate(g["gaps"], 1):
            qhead(doc, i, b["q"]); para(doc, b["answer"])

    if g.get("plan"):
        heading(doc, "30-60-90 day outline")
        for k, lab in (("d30", "First 30 days"), ("d60", "Days 31–60"), ("d90", "Days 61–90")):
            if g["plan"].get(k):
                p = doc.add_paragraph(); spacing(p, 4, 1); p.paragraph_format.keep_with_next = True
                runs(p, lab, bold_all=True)
                for x in g["plan"][k]:
                    bullet(doc, x)

    if g.get("ask_them"):
        heading(doc, "Questions to ask them")
        for x in g["ask_them"]:
            bullet(doc, x)
    if g.get("logistics"):
        heading(doc, "Before and after")
        for x in g["logistics"]:
            bullet(doc, x)

    doc.save(out + ".docx")
    outdir = os.path.dirname(os.path.abspath(out))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, out + ".docx"],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pages = subprocess.run(["pdfinfo", out + ".pdf"], capture_output=True, text=True).stdout
    n = int(re.search(r"Pages:\s+(\d+)", pages).group(1))
    print(f"{out}.docx / .pdf -> {n} page(s)")
    if n > 7:
        print("WARNING: over 7 pages - tighten answers and re-render")


if __name__ == "__main__":
    build(json.load(open(sys.argv[1])), sys.argv[2])
